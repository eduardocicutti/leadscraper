import base64
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import parse_qs, unquote, urlparse

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.api.app import create_app


def make_client(tmp_path: str) -> TestClient:
    os.environ["DB_PATH"] = str(Path(tmp_path) / "test.db")
    return TestClient(create_app())


def csv_payload(text: str, filename: str = "leads.csv") -> dict:
    return {
        "filename": filename,
        "content_base64": base64.b64encode(text.encode("utf-8")).decode("ascii"),
        "prospectador": "Eduardo",
    }


def whatsapp_text(link: str) -> str:
    parsed = urlparse(link)
    text = parse_qs(parsed.query).get("text", [""])[0]
    return unquote(text)


def test_diagnostics_reports_database_counts():
    with TemporaryDirectory() as tmp:
        with make_client(tmp) as client:
            response = client.get("/diagnostics")
            assert response.status_code == 200
            data = response.json()
            assert data["ok"] is True
            assert data["integrity_check"] == "ok"
            assert data["counts"]["selected_leads"] == 0
            assert data["db_path"].endswith("test.db")


def test_spreadsheet_import_deduplicates_by_phone():
    with TemporaryDirectory() as tmp:
        with make_client(tmp) as client:
            content = """Nome da empresa;Telefone/WhatsApp;Cidade;UF;Ramo de atividade
Clínica Boa Vida;(11) 91234-5678;São Paulo;SP;Clínica
Clinica Boa Vida Ltda;11912345678;São Paulo;SP;Clínica
Padaria Sol;(11) 3333-4444;Campinas;SP;Padaria
"""
            response = client.post(
                "/selected-leads/import-spreadsheet",
                json=csv_payload(content),
            )
            assert response.status_code == 200
            data = response.json()
            assert data["imported_count"] == 2
            assert data["skipped_count"] == 1
            assert data["skipped"][0]["reason"] == "telefone"

            selected = client.get("/selected-leads").json()
            assert len(selected) == 2


def test_segment_template_refreshes_only_matching_segment_message():
    with TemporaryDirectory() as tmp:
        with make_client(tmp) as client:
            content = """Nome da empresa;Telefone/WhatsApp;Cidade;UF;Ramo de atividade
Clínica Boa Vida;(11) 91234-5678;São Paulo;SP;Clínica
Academia Norte;(11) 99888-7777;São Paulo;SP;Academia
"""
            client.post("/selected-leads/import-spreadsheet", json=csv_payload(content))

            response = client.put(
                "/selected-leads/segment-templates",
                json={
                    "segmento": "Clínica",
                    "template": "Mensagem clínica para {empresa} em {cidade}",
                },
            )
            assert response.status_code == 200

            selected = client.get("/selected-leads").json()
            by_name = {lead["nome"]: lead for lead in selected}
            clinic_text = whatsapp_text(by_name["Clínica Boa Vida"]["whatsapp_link"])
            academy_text = whatsapp_text(by_name["Academia Norte"]["whatsapp_link"])

            assert "Mensagem clínica" in clinic_text
            assert "Mensagem clínica" not in academy_text


def test_filtered_export_contains_only_matching_selected_leads():
    with TemporaryDirectory() as tmp:
        with make_client(tmp) as client:
            content = """Nome da empresa;Telefone/WhatsApp;Cidade;UF;Ramo de atividade
Clínica Boa Vida;(11) 91234-5678;São Paulo;SP;Clínica
Padaria Sol;(11) 3333-4444;Campinas;SP;Padaria
"""
            client.post("/selected-leads/import-spreadsheet", json=csv_payload(content))

            response = client.post(
                "/selected-leads/export",
                json={
                    "segmento": "Clínica",
                    "cidade": "",
                    "estado": "",
                    "temperatura": "",
                    "com_whatsapp": "",
                    "com_site": "",
                    "prospectador": "Eduardo",
                },
            )
            assert response.status_code == 200
            workbook_path = Path(tmp) / "export.xlsx"
            workbook_path.write_bytes(response.content)
            workbook = load_workbook(workbook_path)
            sheet = workbook.active

            company_names = [sheet.cell(row=row, column=3).value for row in range(2, sheet.max_row + 1)]
            assert "Clínica Boa Vida" in company_names
            assert "Padaria Sol" not in company_names


def test_backup_and_restore_round_trip_database():
    with TemporaryDirectory() as tmp:
        with make_client(tmp) as client:
            content = """Nome da empresa;Telefone/WhatsApp;Cidade;UF;Ramo de atividade
ClÃ­nica Boa Vida;(11) 91234-5678;SÃ£o Paulo;SP;ClÃ­nica
"""
            client.post("/selected-leads/import-spreadsheet", json=csv_payload(content))

            backup = client.get("/backup")
            assert backup.status_code == 200
            assert len(backup.content) > 0

            lead = client.get("/selected-leads").json()[0]
            delete_response = client.delete(f"/selected-leads/{lead['id']}")
            assert delete_response.status_code == 200
            assert client.get("/selected-leads").json() == []

            restore_response = client.post(
                "/restore",
                json={
                    "filename": "backup.db",
                    "content_base64": base64.b64encode(backup.content).decode("ascii"),
                },
            )
            assert restore_response.status_code == 200
            restored = client.get("/selected-leads").json()
            assert len(restored) == 1
            assert restored[0]["nome"] == "ClÃ­nica Boa Vida"
