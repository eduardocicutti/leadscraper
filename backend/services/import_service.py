import csv
import io
from pathlib import Path
from typing import Any
from unicodedata import combining, normalize

from openpyxl import load_workbook

from backend.domain.scoring import classify_porte, score_lead
from backend.domain.whatsapp import build_whatsapp_link, phone_digits


COLUMN_ALIASES = {
    "nome": {
        "nome",
        "empresa",
        "nome da empresa",
        "razao social",
        "razão social",
        "company",
        "company name",
    },
    "categoria": {
        "categoria",
        "segmento",
        "ramo",
        "ramo de atividade",
        "atividade",
        "tipo",
    },
    "telefone": {
        "telefone",
        "telefone/whatsapp",
        "whatsapp",
        "celular",
        "phone",
        "numero",
        "número",
    },
    "site": {"site", "website", "url", "link site"},
    "url_maps": {
        "maps",
        "google maps",
        "link maps",
        "url maps",
        "mapa",
        "onde encontrou?",
    },
    "cidade": {"cidade", "city", "municipio", "município"},
    "estado": {"estado", "uf", "state"},
    "endereco": {"endereco", "endereço", "address", "localizacao", "localização"},
    "nota": {"nota", "nota google", "rating", "avaliação", "avaliacao"},
    "avaliacoes": {"avaliacoes", "avaliações", "reviews", "qtd avaliações", "aval."},
    "porte": {"porte", "tamanho"},
    "classificacao": {"classificacao", "classificação", "temperatura", "temp."},
    "score": {"score", "pontuacao", "pontuação"},
    "prospectador": {"prospectador", "responsavel", "responsável", "owner"},
    "notes": {"notas", "observacoes", "observações", "notes"},
}


def _normalize_header(value: Any) -> str:
    text_value = normalize("NFKD", str(value or "").strip().lower())
    ascii_value = "".join(ch for ch in text_value if not combining(ch))
    return " ".join("".join(ch if ch.isalnum() else " " for ch in ascii_value).split())


ALIAS_TO_FIELD = {
    _normalize_header(alias): field
    for field, aliases in COLUMN_ALIASES.items()
    for alias in aliases
}


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_float(value: Any) -> float | None:
    text = _clean(value).replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _to_int(value: Any) -> int | None:
    text = "".join(ch for ch in _clean(value) if ch.isdigit())
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        return None


def _is_mobile_phone(value: str) -> bool:
    digits = phone_digits(value)
    if digits.startswith("55") and len(digits) in (12, 13):
        digits = digits[2:]
    return len(digits) == 11 and digits[2] == "9"


def _headers_to_fields(headers: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for index, header in enumerate(headers):
        field = ALIAS_TO_FIELD.get(_normalize_header(header))
        if field:
            mapping[index] = field
    return mapping


def _row_to_lead(row: list[Any], mapping: dict[int, str], fallback_prospectador: str) -> dict:
    item: dict[str, Any] = {}
    for index, value in enumerate(row):
        field = mapping.get(index)
        if not field:
            continue
        item[field] = _clean(value)

    telefone = item.get("telefone") or ""
    categoria = item.get("categoria") or item.get("segmento") or ""
    nota = _to_float(item.get("nota"))
    avaliacoes = _to_int(item.get("avaliacoes"))
    score = _to_int(item.get("score")) or 0
    porte = item.get("porte") or classify_porte(item.get("nome", ""), categoria, avaliacoes or 0)

    lead = {
        "nome": item.get("nome", ""),
        "categoria": categoria,
        "segmento": categoria,
        "nota": nota,
        "avaliacoes": avaliacoes,
        "endereco": item.get("endereco") or None,
        "telefone": telefone or None,
        "is_whatsapp": _is_mobile_phone(telefone),
        "whatsapp_link": "",
        "site": item.get("site") or None,
        "url_maps": item.get("url_maps") or "",
        "porte": porte,
        "classificacao": item.get("classificacao") or "",
        "score": score,
        "cidade": item.get("cidade") or "",
        "estado": item.get("estado") or "",
        "prospectador": item.get("prospectador") or fallback_prospectador,
        "notes": item.get("notes") or "",
    }
    if not lead["score"] or not lead["classificacao"]:
        classificacao, calculated_score = score_lead(lead)
        lead["classificacao"] = lead["classificacao"] or classificacao
        lead["score"] = lead["score"] or calculated_score
    if lead["is_whatsapp"]:
        lead["whatsapp_link"] = build_whatsapp_link(telefone, lead["nome"])
    return lead


def _rows_from_csv(content: bytes, suffix: str) -> list[list[Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    delimiter = "\t" if suffix == ".tsv" else None
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;|\t").delimiter
        except csv.Error:
            delimiter = ";"
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _rows_from_workbook(content: bytes) -> list[list[Any]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def parse_leads_spreadsheet(
    filename: str,
    content: bytes,
    fallback_prospectador: str = "",
) -> tuple[list[dict], dict]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".csv", ".tsv"}:
        rows = _rows_from_csv(content, suffix)
    elif suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        rows = _rows_from_workbook(content)
    else:
        raise ValueError("Formato não suportado. Use .xlsx, .csv ou .tsv.")

    rows = [row for row in rows if any(_clean(cell) for cell in row)]
    if not rows:
        return [], {"rows_read": 0, "columns_detected": []}

    header_index = 0
    mapping: dict[int, str] = {}
    for index, row in enumerate(rows[:10]):
        candidate = _headers_to_fields(row)
        if len(candidate) > len(mapping):
            header_index = index
            mapping = candidate
    if not mapping:
        raise ValueError("Não encontrei cabeçalhos reconhecíveis na planilha.")

    leads = [
        _row_to_lead(row, mapping, fallback_prospectador)
        for row in rows[header_index + 1 :]
    ]
    leads = [lead for lead in leads if lead.get("nome") or lead.get("telefone")]
    return leads, {
        "rows_read": max(len(rows) - header_index - 1, 0),
        "columns_detected": sorted(set(mapping.values())),
    }
